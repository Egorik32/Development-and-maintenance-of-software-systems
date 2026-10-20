import json
import re
import sys
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
LABS = json.loads((Path(__file__).parent / "lab_conditions.json").read_text(encoding="utf-8"))
PKG = "src/main/java/com/example/cityparks/"
RES = "src/main/resources/com/example/cityparks/"

WORK = {
    3: "Создан проект JavaFX с пакетами для модели, представлений FXML и контроллеров. Разметка загружена, проверена и сохранена через модель редактора Scene Builder. Главный класс загружает RootLayout и CityOverview. В этом этапе интерфейс еще не содержит данных и обработчиков изменения записей.",
    4: "Созданы модели City и Park с JavaFX Properties. Тестовые данные помещены в ObservableList, а контроллер связывает свойства названия и статуса с колонками TableView.",
    5: "Добавлены обработчики выбора строки, добавления, редактирования и удаления города. Диалог содержит поля города и парка из варианта 12. Запись изменяется только после проверки ввода и нажатия OK. Отмена сохраняет исходные значения.",
    6: "К основной сцене и диалогу подключена таблица стилей CSS. В заголовок основного окна и диалога добавлена иконка парка. Изменены оформление меню, таблицы и заголовков.",
    7: "Созданы обертка CityListWrapper и адаптер LocalDateAdapter. JAXB сохраняет и загружает XML. Контроллер меню использует FileChooser для открытия и сохранения файла. Preferences хранит абсолютный путь к последнему файлу; при следующем запуске этот файл загружается автоматически.",
    8: "Нативный пакет создан генератором AntTemplate из официального плагина e(fx)clipse 2.4.0. Сохранены build.fxbuild, исходный экспорт build.generated.xml и рабочий build.xml. Ant сформировал CityParks.exe и собственную среду выполнения Java. Скрипт запуска переносит пакет в каталог без кириллицы, необходимый для старого Windows-лаунчера JavaFX 8. Дополнительно по запросу добавлен пункт меню «Вид» → «Тёмная тема». CheckMenuItem связан с BooleanProperty; Preferences сохраняет выбор. Светлый CSS подключается всегда, тёмный добавляется при включении режима. Тема применяется к главному окну, редактированию и сообщениям об ошибках.",
    9: "Создан проект Dynamic Web Project с уровнем Servlet 5.0 для Apache Tomcat 10.0. Главная страница index.jsp содержит меню городов и парков. Общие блоки header.jsp и footer.jsp подключаются через jsp:include. Локально подключены jQuery 3.6.4 и Bootstrap 5 с Popper. При наведении на тестовую кнопку появляется сообщение «Ура, Bootstrap 5 работает»; инициализация выполняется через jQuery. Подготовлены файлы проекта Eclipse и локальный репозиторий Git.",
    10: "Созданы шаблоны CityServlet и ParkServlet в пакете controller. Аннотации WebServlet определяют адреса. doGet формирует HTML-ответ UTF-8, doPost вызывает doGet. Обе страницы проверены на работающем Tomcat.",
    11: "Создан пакет domain с сущностями City и Park. Поля соответствуют варианту 12. Добавлены технические идентификаторы, внешний ключ idPark и навигационное свойство park по образцу связи сущностей в методичке. Определены конструкторы, геттеры, сеттеры и toString.",
    12: "Разработаны представления city.jsp и park.jsp. Сервлеты передают тестовые коллекции моделей в request. JSTL forEach отображает строки таблиц, список парков заполняет селектор формы города. Кнопка добавления на данном этапе отключена; ее обработчик реализуется в работе 1.14.",
    13: "Создана отдельная база PostgreSQL cityparks_lab13 с таблицами parks и cities и внешним ключом. Реализованы ConnectionProperty, PostgreSQLConnectionBuilder и классы DAO. Сервлеты получают записи из базы через findAll и передают их в JSP.",
    14: "В отдельной базе cityparks_lab14 реализовано добавление парков и городов. doPost читает параметры формы и вызывает insert. PreparedStatement выполняет INSERT, а база генерирует первичный ключ. Город сохраняется со ссылкой на выбранный парк.",
    15: "В отдельной базе cityparks_lab15 реализованы изменение и удаление городов и парков. EditCityServlet и EditParkServlet загружают выбранную запись, а doPost вызывает update. DeleteCityServlet и DeleteParkServlet вызывают delete. Перед переходом по ссылке удаления браузер запрашивает подтверждение.",
}

LISTINGS = {
    3: [(PKG + "MainApp.java", "initRootLayout"), (RES + "view/RootLayout.fxml", None)],
    4: [(PKG + "MainApp.java", "MainApp"), (PKG + "controller/CityOverviewController.java", "initialize"), (PKG + "model/City.java", "nameProperty")],
    5: [(PKG + "controller/CityOverviewController.java", "handleNewCity"), (PKG + "controller/CityEditDialogController.java", "isInputValid")],
    6: [(PKG + "MainApp.java", "start"), (RES + "css/style.css", None)],
    7: [(PKG + "model/CityListWrapper.java", None), (PKG + "util/XmlDataStore.java", "save"), (PKG + "MainApp.java", "setCityFilePath")],
    8: [("build.fxbuild", None), ("run.ps1", None), (PKG + "controller/RootLayoutController.java", "setMainApp"), (PKG + "MainApp.java", "applyTheme")],
    9: [("src/main/webapp/index.jsp", None), ("src/main/webapp/js/bootstrap-check.js", None)],
    10: [("src/main/java/controller/CityServlet.java", None), ("src/main/java/controller/ParkServlet.java", "doPost")],
    11: [("src/main/java/domain/City.java", "fields"), ("src/main/java/domain/City.java", "City"), ("src/main/java/domain/City.java", "toString"), ("src/main/java/domain/Park.java", "fields"), ("src/main/java/domain/Park.java", "Park")],
    12: [("src/main/java/controller/CityServlet.java", "doGet"), ("src/main/webapp/city.jsp", "JSTL")],
    13: [("src/main/java/dao/PostgreSQLConnectionBuilder.java", None), ("src/main/java/dao/CityDbDAO.java", "findAll")],
    14: [("src/main/java/controller/CityServlet.java", "doPost"), ("src/main/java/dao/ParkDbDAO.java", "insert")],
    15: [("src/main/java/controller/EditCityServlet.java", "doPost"), ("src/main/java/dao/CityDbDAO.java", "update"), ("src/main/java/dao/ParkDbDAO.java", "delete")],
}


def excerpt(folder, filename, method):
    text = (folder / filename).read_text(encoding="utf-8")
    if method == 'fields':
        start = text.index('    private ')
        end = text.index('    public ', start)
        return text[start:end].strip()
    if method == 'JSTL':
        start = text.index('<c:forEach')
        end = text.index('</c:forEach>', start) + len('</c:forEach>')
        return text[start:end]
    if method:
        parameters = r'\([^)]*String[^)]*\)' if method in ('City', 'Park') else r'\([^)]*\)'
        match = re.search(r'(?m)^\s*(?:@Override\s*)?(?:public|private|protected|static)[^\n]*\b' + method + parameters + r'\s*(?:throws [^{]+)?\{', text)
        if not match: raise ValueError(filename + ": " + method)
        begin = match.start()
        end = text.index('{', begin) + 1
        depth = 1
        while depth:
            depth += (text[end] == '{') - (text[end] == '}')
            end += 1
        text = text[begin:end].strip()
    if filename.endswith('.css'):
        text = '\n'.join(text.splitlines()[:35])
    return text


def code(doc, text):
    lines = text.splitlines()
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = len(lines) <= 12 and index < len(lines) - 1
        run = p.add_run(line or ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(8)


def make_report(lab):
    folder = ROOT / lab['folder']
    n = lab['number']
    result_path = folder / 'verification/result.txt'
    if not result_path.is_file() or 'ALL CHECKS PASSED' not in result_path.read_text(encoding='utf-8-sig'):
        raise ValueError('Missing successful verification: ' + lab['folder'])
    if n == 8 and not (folder / 'verification/native.txt').exists():
        raise ValueError('Native executable verification missing')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(1.6)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if key.endswith('Theme'): del fonts.attrib[key]
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'): fonts.set(qn('w:' + key), 'Times New Roman')
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(16)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 2'].font.size = Pt(11)
    p = doc.add_paragraph(f'Отчет по лабораторной работе 1 {n}', style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(lab['title'])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    doc.add_paragraph('Разработка и сопровождение программных систем. Вариант 12. Группа ПИ-331.')
    doc.add_heading('Цель работы', 1)
    doc.add_paragraph(lab['title'] + ' для предметной области «Города и парки».')
    doc.add_heading('Условие работы', 1)
    task = (folder / 'ЗАДАНИЕ.txt').read_text(encoding='utf-8')
    if n < 9: task = lab['condition']
    else: task = task.split('Задание на лабораторную работу', 1)[-1]
    task = re.sub(r'(?m)^\s*\d{2}\s*$', '', task)
    if n < 9:
        for block in task.split('\n\n'): doc.add_paragraph(block.strip())
    else:
        for block in re.split(r'(?m)(?=^\s*\d+\.\s)', task):
            normalized = re.sub(r'\s+', ' ', block).strip()
            if normalized: doc.add_paragraph(normalized)
    doc.add_heading('Вариант и модель данных', 1)
    doc.add_paragraph('Города: название, статус, население. Парки: название, тип, площадь, наличие водоема, дата постройки.')
    doc.add_heading('Ход работы', 1)
    doc.add_paragraph(WORK[n])
    if n >= 9:
        doc.add_paragraph(f'Код этапа сохранен в локальной ветке dev{n - 8}; рабочий результат доступен в master. История предыдущих этапов сохранена в репозитории этой папки.')
    doc.add_heading('Фрагменты исходного кода', 1)
    doc.add_paragraph('Ниже приведены фрагменты файлов проекта. Полные исходники находятся в папке лабораторной работы.')
    for filename, method in LISTINGS[n]:
        label = ' поля модели' if method == 'fields' else (' цикл JSTL' if method == 'JSTL' else (f' метод {method}' if method else ''))
        doc.add_heading(Path(filename).stem + ' ' + Path(filename).suffix.lstrip('.') + label, 2)
        code(doc, excerpt(folder, filename, method))
    doc.add_heading('Проверка работы', 1)
    result = result_path.read_text(encoding='utf-8-sig')
    checks = [line.removeprefix('PASS: ') for line in result.splitlines() if line.startswith('PASS: ')]
    if n == 8:
        checks = ['загрузка FXML и иконки', 'оба режима оформления и синхронизация переключателя', 'сохранение темы после перезапуска', 'тёмная тема диалога и сообщений об ошибках', 'добавление, изменение, отмена и удаление', 'блокирование некорректного ввода', 'сохранение и загрузка XML', 'автоматическое открытие последнего XML', 'сохранение данных при ошибочном XML']
    verified_date = datetime.fromtimestamp(result_path.stat().st_mtime).strftime('%d.%m.%Y')
    doc.add_paragraph('Локальная проверка выполнена ' + verified_date + '. Сборка завершилась без ошибок. Подтверждены: ' + '; '.join(dict.fromkeys(checks)) + '.')
    if n == 8:
        doc.add_paragraph('Дополнительно проверен запуск CityParks.exe с собственной средой выполнения: журнал подтверждает открытое окно JavaFX и загруженное центральное представление.')
        doc.add_paragraph('Скрипт запуска ожидает закрытия графического приложения и проверяет его ExitCode. Код успешного копирования robocopy больше не принимается за ошибку EXE.')
        doc.add_paragraph('Светлая и тёмная темы проверены также в JAR нативного пакета. Переключение не изменяет записи; после повторного создания окна восстанавливается выбор темы. Сохранение XML и проверка пользовательского ввода работают в тёмном режиме.')
    image_path = folder / 'verification/main.png'
    if n == 9: image_path = folder / 'verification/bootstrap-desktop.png'
    if image_path.exists():
        doc.add_heading('Результат запуска', 1)
        p = doc.add_paragraph()
        p.add_run().add_picture(str(image_path), width=Cm(16.5))
        doc.add_paragraph('Главная страница с проверочным сообщением Bootstrap.' if n == 9 else 'Окно приложения при локальной проверке.')
    doc.add_heading('Запуск', 1)
    doc.add_paragraph('В VS Code открыть папку этой лабораторной и терминал PowerShell. Выполнить .\\run.ps1.' + (f' Перед первым запуском выполнить .\\setup_db.ps1 для подготовки тестовой базы cityparks_lab{n}.' if n >= 13 else ''))
    if n >= 9:
        doc.add_paragraph(f'Адрес приложения: http://localhost:8080/lab1-{n}/. Перед запуском следующей веб-работы остановить текущий сервер через Ctrl+C.')
    if n == 9:
        doc.add_paragraph('Для проверки библиотек навести курсор на кнопку «Поднеси ко мне курсор». Появляется сообщение «Ура, Bootstrap 5 работает», которое скрывается после ухода курсора. Библиотеки загружаются из файлов проекта без обращения к CDN.')
    if n == 8:
        doc.add_paragraph('В меню «Вид» выбрать «Тёмная тема», открыть редактирование и перезапустить приложение. Выбор должен сохраниться; снятие галочки возвращает светлый режим. Запуск run.ps1 пересобирает и обновляет нативную копию.')
    if n == 11:
        doc.add_paragraph('Для демонстрации конструкторов, полей, связи сущностей и методов toString выполнить .\\show_domain.ps1.')
    doc.add_heading('Вывод', 1)
    doc.add_paragraph('Результат программной части этого этапа получен и проверен. Подготовлены исходники, сценарий запуска и отчет для демонстрации преподавателю.')
    doc.add_heading('Источники', 1)
    doc.add_paragraph(f'https://do.rsue.ru/mod/assign/view.php?id={129763+n}')
    if 'tutorial' in lab: doc.add_paragraph(lab['tutorial'])
    else: doc.add_paragraph('Методические указания Moodle, оригинал УСЛОВИЕ_ОРИГИНАЛ.pdf в папке работы.')
    output = folder / f'Отчет_ЛР_1_{n}_вариант_12.docx'
    doc.save(output)
    print(output, flush=True)


if __name__ == '__main__':
    for lab in LABS:
        if len(sys.argv) == 1 or str(lab['number']) in sys.argv[1:]: make_report(lab)
