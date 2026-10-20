from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)


def set_base_style(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.4)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = None
    styles["Title"].font.size = Pt(16)
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 2"].font.size = Pt(12)


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.add_run(subtitle).bold = True


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Раздел"
    table.rows[0].cells[1].text = "Содержание"
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    doc.add_paragraph()


def add_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def read_text(path):
    return (ROOT / path).read_text(encoding="utf-8")


def excerpt(path, start=None, end=None):
    lines = read_text(path).splitlines()
    if start is None and end is None:
        return "\n".join(lines)
    start_index = 0 if start is None else max(start - 1, 0)
    end_index = len(lines) if end is None else min(end, len(lines))
    return "\n".join(lines[start_index:end_index])


def add_code_listing(doc, title, code):
    doc.add_heading(title, level=2)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(32, 32, 32)


def add_code_section(doc, listings):
    if not listings:
        return
    doc.add_heading("Листинги кода", level=1)
    for title, code in listings:
        add_code_listing(doc, title, code)


def save_report(filename, title, lab, purpose, implemented, files, verification, listings=None):
    doc = Document()
    set_base_style(doc)
    add_title(doc, title, f"{lab}. Вариант 12")
    doc.add_heading("Цель работы", level=1)
    doc.add_paragraph(purpose)
    doc.add_heading("Предметная область", level=1)
    add_kv_table(doc, [
        ("Таблица 1", "Города: название, статус, население."),
        ("Таблица 2", "Парки: название, тип, площадь, наличие водоема, дата постройки."),
    ])
    doc.add_heading("Выполненная работа", level=1)
    add_list(doc, implemented)
    doc.add_heading("Структура решения", level=1)
    add_list(doc, files)
    add_code_section(doc, listings or [])
    doc.add_heading("Проверка", level=1)
    add_list(doc, verification)
    doc.add_heading("Вывод", level=1)
    doc.add_paragraph(
        "Лабораторная работа выполнена по варианту 12. Реализация соответствует заданной "
        "предметной области и проверена локальной сборкой."
    )
    path = REPORTS / filename
    doc.save(path)
    return path


created = [
    save_report(
        "Отчет_ЛР_1_1_Анонимные_классы_вариант_12.docx",
        "Отчет по лабораторной работе",
        "ЛР 1.1 Анонимные классы",
        "Создать Java-приложение с использованием анонимных классов.",
        [
            "Созданы классы предметной области City и Park.",
            "Определены интерфейсы Checker и Formatter.",
            "Через анонимные классы реализованы проверки крупного города и наличия водоема в парке.",
            "Через анонимный класс реализовано форматирование строк вывода.",
        ],
        [
            "Lab1_AnonymousClasses_V12/src/Lab1AnonymousClassesV12.java",
            "Lab1_AnonymousClasses_V12/README.md",
        ],
        [
            "Исходный файл скомпилирован командой javac с кодировкой UTF-8.",
            "Программа запущена командой java -cp out Lab1AnonymousClassesV12.",
            "Вывод содержит данные Ростова-на-Дону, Сочи и Шахт, а также сведения о парках.",
        ],
        [
            (
                "Листинг 1 Основной класс с анонимными классами",
                excerpt("Lab1_AnonymousClasses_V12/src/Lab1AnonymousClassesV12.java"),
            ),
        ],
    ),
    save_report(
        "Отчет_ЛР_1_2_Лямбда_выражения_вариант_12.docx",
        "Отчет по лабораторной работе",
        "ЛР 1.2 Лямбда выражения",
        "Создать Java-приложение с использованием лямбда-выражений и функциональных интерфейсов.",
        [
            "Созданы классы City и Park для варианта 12.",
            "Использованы Predicate, Function и собственный функциональный интерфейс Metric.",
            "Реализованы фильтрация городов с парками у воды, сортировка, расчет среднего населения.",
            "Добавлен расчет нагрузки населения на площадь парка.",
        ],
        [
            "Lab2_LambdaExpressions_V12/src/Lab2LambdaExpressionsV12.java",
            "Lab2_LambdaExpressions_V12/README.md",
        ],
        [
            "Исходный файл скомпилирован командой javac с кодировкой UTF-8.",
            "Программа запущена командой java -cp out Lab2LambdaExpressionsV12.",
            "Результат показывает фильтрацию, сортировку, среднее значение и минимальную дату постройки парка.",
        ],
        [
            (
                "Листинг 1 Основной класс с лямбда выражениями",
                excerpt("Lab2_LambdaExpressions_V12/src/Lab2LambdaExpressionsV12.java"),
            ),
        ],
    ),
    save_report(
        "Отчет_ЛР_1_3_1_8_JavaFX_вариант_12.docx",
        "Отчет по лабораторным работам",
        "ЛР 1.3-1.8 JavaFX приложение",
        "Разработать JavaFX-приложение с интерфейсом, моделью данных, таблицей, операциями редактирования, стилями и XML-хранением.",
        [
            "Создан JavaFX-проект CityParksJavaFX для варианта 12.",
            "Реализована модель City и Park на JavaFX properties.",
            "Созданы FXML-представления главного окна, таблицы городов и диалога редактирования.",
            "Добавлены операции создания, редактирования и удаления городов.",
            "Настроены CSS-стили и работа с XML-файлом cities.xml.",
            "Сохранение последнего пути к файлу выполнено через Preferences.",
        ],
        [
            "CityParksJavaFX/src/main/java/com/example/cityparks/MainApp.java",
            "CityParksJavaFX/src/main/java/com/example/cityparks/model/City.java",
            "CityParksJavaFX/src/main/java/com/example/cityparks/model/Park.java",
            "CityParksJavaFX/src/main/java/com/example/cityparks/controller",
            "CityParksJavaFX/src/main/resources/com/example/cityparks/view",
            "CityParksJavaFX/src/main/resources/com/example/cityparks/css/style.css",
            "CityParksJavaFX/data/cities.xml",
        ],
        [
            "Проект скомпилирован скриптом CityParksJavaFX/build.ps1.",
            "Сборка JavaFX завершилась без ошибок.",
            "Проверено отсутствие старой предметной области в собственных файлах проекта.",
        ],
        [
            (
                "Листинг 1 Модель города",
                excerpt("CityParksJavaFX/src/main/java/com/example/cityparks/model/City.java"),
            ),
            (
                "Листинг 2 Фрагмент контроллера таблицы",
                excerpt("CityParksJavaFX/src/main/java/com/example/cityparks/controller/CityOverviewController.java", 1, 115),
            ),
            (
                "Листинг 3 Фрагмент FXML представления",
                excerpt("CityParksJavaFX/src/main/resources/com/example/cityparks/view/CityOverview.fxml", 1, 90),
            ),
            (
                "Листинг 4 Фрагмент XML хранилища",
                excerpt("CityParksJavaFX/src/main/java/com/example/cityparks/util/XmlDataStore.java", 1, 95),
            ),
        ],
    ),
    save_report(
        "Отчет_ЛР_1_9_1_15_JavaEE_вариант_12.docx",
        "Отчет по лабораторным работам",
        "ЛР 1.9-1.15 Java EE веб приложение",
        "Разработать веб-приложение Dynamic Web Project для вывода, создания, редактирования и удаления данных из PostgreSQL.",
        [
            "Создан веб-проект CityParksJavaEE для Tomcat 10 и Jakarta Servlet.",
            "Разработана главная страница index.jsp и страницы city.jsp, park.jsp, editcity.jsp, editpark.jsp.",
            "Созданы доменные классы City и Park.",
            "Реализованы сервлеты списка, добавления, редактирования и удаления.",
            "Разработаны DAO-классы для PostgreSQL с PreparedStatement.",
            "Подготовлен SQL-скрипт создания таблиц cities и parks с тестовыми данными.",
            "В проект добавлены Bootstrap, JSTL, Servlet API и PostgreSQL JDBC Driver.",
        ],
        [
            "CityParksJavaEE/src/main/java/domain",
            "CityParksJavaEE/src/main/java/dao",
            "CityParksJavaEE/src/main/java/controller",
            "CityParksJavaEE/src/main/webapp",
            "CityParksJavaEE/sql/schema.sql",
            "CityParksJavaEE/build.ps1",
        ],
        [
            "Java-классы веб-проекта скомпилированы скриптом CityParksJavaEE/build.ps1.",
            "Сборка DAO, домена и сервлетов завершилась без ошибок.",
            "JSP-страницы подготовлены для проверки при запуске проекта на Tomcat.",
        ],
        [
            (
                "Листинг 1 Доменный класс City",
                excerpt("CityParksJavaEE/src/main/java/domain/City.java"),
            ),
            (
                "Листинг 2 DAO для городов",
                excerpt("CityParksJavaEE/src/main/java/dao/CityDbDAO.java", 1, 100),
            ),
            (
                "Листинг 3 Сервлет городов",
                excerpt("CityParksJavaEE/src/main/java/controller/CityServlet.java"),
            ),
            (
                "Листинг 4 Фрагмент JSP страницы городов",
                excerpt("CityParksJavaEE/src/main/webapp/city.jsp", 1, 95),
            ),
            (
                "Листинг 5 SQL схема базы данных",
                excerpt("CityParksJavaEE/sql/schema.sql"),
            ),
        ],
    ),
]

for path in created:
    print(path)
