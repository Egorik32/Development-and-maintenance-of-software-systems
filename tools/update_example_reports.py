import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from create_individual_reports import excerpt

ROOT = Path(__file__).resolve().parents[1]
TOPICS = [
    "Сложение через лямбду", "Эквивалентный анонимный класс",
    "Арифметические операции", "Терминальная лямбда",
    "Переменные класса", "Локальные переменные", "Блочная лямбда",
    "Обобщенный функциональный интерфейс", "Лямбда как параметр",
    "Ссылки на статические методы", "Ссылка на метод объекта",
    "Ссылка на конструктор", "Лямбда как результат метода",
]


def insert_paragraph(doc, anchor, text, style=None):
    paragraph = doc.add_paragraph(text, style)
    paragraph.paragraph_format.space_after = Pt(5)
    anchor.addprevious(paragraph._p)
    return paragraph


def insert_heading(doc, anchor, text, level):
    paragraph = doc.add_heading(text, level)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    anchor.addprevious(paragraph._p)


def insert_code(doc, anchor, text):
    for line in text.splitlines():
        paragraph = insert_paragraph(doc, anchor, "")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = False
        run = paragraph.add_run(line or " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def results_table(doc, anchor, manifest, results):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(0.9), Cm(11.5), Cm(4.4)]
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, title in zip(table.rows[0].cells, ["№", "Учебный пример", "Вывод программы"]):
        cell.text = title
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for example, result in zip(manifest, results):
        row = table.add_row()
        for cell, value in zip(row.cells, [str(example["number"]), TOPICS[example["number"] - 1], "; ".join(result["output"])]):
            cell.text = value
    for index, row in enumerate(table.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + side)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
            props.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7E7E7" if index == 0 else ("FFFFFF" if index % 2 else "F6F6F6"))
            props.append(shading)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "70")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    anchor.addprevious(table._tbl)


def update(number):
    folder = ROOT / ("Lab1_AnonymousClasses_V12" if number == 1 else "Lab2_LambdaExpressions_V12")
    filename = "Отчет_ЛР_1_1_Анонимные_классы_вариант_12.docx" if number == 1 else "Отчет_ЛР_1_2_Лямбда_выражения_вариант_12.docx"
    output = ROOT / "reports" / filename
    backup = folder / "verification" / "report_before_examples.docx"
    if not backup.exists():
        shutil.copy2(output, backup)
    doc = Document(output)
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    anchor = next(p._p for p in doc.paragraphs if p.text == "Проверка" and p.style.name == "Heading 1")
    # Replace only the previously inserted section on a repeated update.
    begin = next((p._p for p in doc.paragraphs if p.text == "Учебные примеры из методички"), None)
    while begin is not None and begin is not anchor:
        following = begin.getnext()
        begin.getparent().remove(begin)
        begin = following
    manifest = json.loads((folder / "examples/examples.json").read_text(encoding="utf-8"))
    results = json.loads((folder / "verification/examples_result.json").read_text(encoding="utf-8-sig"))
    if len(manifest) != len(results) or not all(result["passed"] for result in results):
        raise ValueError("The complete example verification is required")
    insert_heading(doc, anchor, "Учебные примеры из методички", 1)
    insert_paragraph(doc, anchor, "Учебные примеры выполняются отдельно от приложения по варианту 12. Исходники находятся в папке examples этой лабораторной. Для каждого примера выполнены компиляция, запуск и сравнение полного вывода с ожидаемым результатом.")
    if number == 1:
        insert_paragraph(doc, anchor, "Пример HelloWorldAnonymous использует интерфейс HelloWorld, обычную реализацию EnglishGreeting и две анонимные реализации для французского и испанского приветствий.")
        for filename in ("HelloWorld.java", "EnglishGreeting.java", "HelloWorldAnonymous.java"):
            insert_heading(doc, anchor, filename.replace(".", " "), 2)
            insert_code(doc, anchor, (folder / "examples/01_HelloWorld" / filename).read_text(encoding="utf-8"))
        insert_paragraph(doc, anchor, "Проверенный вывод: Hello world; Salut Fred; Hola, mundo.")
    else:
        insert_paragraph(doc, anchor, "Одинаковые имена LambdaApp и интерфейсов сохранены в отдельных папках примеров. Каждый пример имеет собственный каталог скомпилированных классов, поэтому имена не конфликтуют.")
        results_table(doc, anchor, manifest, results)
        insert_paragraph(doc, anchor, "Ниже приведены фрагменты методов из файлов examples. Полные классы и необходимые интерфейсы находятся рядом с ними в соответствующей папке.")
        for example in manifest:
            filename = "examples/" + example["folder"] + "/" + example["class"] + ".java"
            insert_heading(doc, anchor, f"Пример {example['number']} {TOPICS[example['number'] - 1]}", 2)
            insert_paragraph(doc, anchor, "Исходный файл: " + filename)
            insert_code(doc, anchor, excerpt(folder, filename, "main"))
            if example["number"] == 9:
                insert_code(doc, anchor, excerpt(folder, filename, "sum"))
            if example["number"] == 13:
                insert_code(doc, anchor, excerpt(folder, filename, "action"))
    insert_heading(doc, anchor, "Запуск и результаты проверки примеров", 1)
    insert_paragraph(doc, anchor, "Открыть папку лабораторной в VS Code и выполнить .\\run_examples.ps1. Для варианта 12 остается команда .\\run.ps1.")
    if number == 2:
        insert_paragraph(doc, anchor, "Один пример запускается по номеру: .\\run_examples.ps1 -Number 12. В этом случае проверяется ссылка на конструктор, вывод программы: Tom.")
    verification_text = "проверен учебный пример" if number == 1 else "проверены все 13 учебных примеров"
    insert_paragraph(doc, anchor, f"14 октября 2026 года {verification_text}. Запуск выполнен в PowerShell 7 и Windows PowerShell 5.1. Результаты находятся в verification/examples_result.json; повторный запуск варианта 12 сохранен в verification/variant_output.txt.")
    for p in doc.paragraphs:
        if p.text == "Вывод":
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        if p.text.startswith(("Лабораторная работа выполнена по варианту 12.", "Программная часть содержит")):
            p.text = "Учебные примеры и приложение по варианту 12 скомпилированы и проверены."
    doc.save(output)
    print(output, flush=True)


if __name__ == "__main__":
    update(1)
    update(2)
